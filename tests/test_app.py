"""Smoke tests: every page renders, and the review and export flow works."""
import io

import openpyxl
import pytest
from streamlit.testing.v1 import AppTest

from tallyhound import common as C

APP = str(__import__("pathlib").Path(__file__).resolve().parent.parent / "app.py")


@pytest.mark.parametrize("page", C.NAV)
def test_every_page_renders(page):
    at = AppTest.from_file(APP, default_timeout=60).run()
    at.radio(key="nav").set_value(page).run()
    assert not at.exception, [e.value for e in at.exception]


def test_review_decisions_and_download_step():
    at = AppTest.from_file(APP, default_timeout=60).run()
    at.session_state.step = 3
    at.run()
    at.button(key="appr_F-01").click().run()
    assert at.session_state.decisions["F-01"]["status"] == "Approved"
    at.session_state.step = 4
    at.run()
    assert not at.exception
    metrics = {m.label: m.value for m in at.metric}
    assert metrics["Approved"] == "1" and metrics["Pending"] == "25"


def test_reject_needs_a_reason():
    at = AppTest.from_file(APP, default_timeout=60).run()
    at.session_state.step = 3
    at.run()
    assert at.button(key="rejbtn_F-02").disabled
    at.text_input(key="rej_F-02").set_value("Not enough evidence").run()
    assert not at.button(key="rejbtn_F-02").disabled


def _exports_app():
    import streamlit as st
    from tallyhound import common as C
    from tallyhound import exports
    C.init_state()
    C.decide("F-01", "Approved")
    C.decide("F-03", "Rejected", "one job")
    st.session_state.wb = exports.build_workbook(False)
    st.session_state.draft = exports.build_workbook(True)
    st.session_state.pdf = exports.build_memo()


def test_exports_build():
    at = AppTest.from_function(_exports_app, default_timeout=60).run()
    assert not at.exception, [e.value for e in at.exception]
    wb = openpyxl.load_workbook(io.BytesIO(at.session_state.wb))
    assert wb.sheetnames == ["Summary", "Monthly audit", "Payment gate", "Supplier recovery", "Subscriptions", "Audit trail"]
    summary = {r[0].value: r[1].value for r in wb["Summary"].iter_rows(min_row=2)}
    assert summary["Approved"] == 1 and summary["Rejected"] == 1 and summary["Pending"] == 24
    draft = openpyxl.load_workbook(io.BytesIO(at.session_state.draft))
    assert "DRAFT" in draft["Summary"]["B2"].value
    assert at.session_state.pdf.startswith(b"%PDF")
